"""Word, Excel and CSV exporters."""

from __future__ import annotations

import csv
from pathlib import Path

from ..core.transcript import Gap, Transcript
from ..i18n import tr_in
from .common import (ExportOptions, Row, any_marked, clock, gap_text, marked_runs, metadata, sentence_rows, sentences,
                     with_gaps)

# Distinct, print-friendly speaker colours (hex RGB).
_SPEAKER_COLORS = ["1F4E79", "B4501E", "2E6B30", "6B3FA0", "8C6A00", "0E7490", "A1263A", "4B5563"]


def _add_page_number(paragraph, template: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    before, _, after = template.partition("{n}")
    if before:
        paragraph.add_run(before)
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    if after:
        paragraph.add_run(after)


def export_docx(t: Transcript, path: Path, options: ExportOptions) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.paragraph_format.space_after = Pt(8)

    doc.core_properties.title = t.source_name
    doc.core_properties.comments = tr_in(options.lang, "exp_generated")
    doc.add_heading(t.source_name, level=1)

    gray = RGBColor(0x6B, 0x72, 0x80)
    meta = doc.add_paragraph()
    for i, (key, value) in enumerate(metadata(t, options.lang)):
        if i:
            meta.add_run("    ")
        label = meta.add_run(f"{key}: ")
        label.bold = True
        label.font.size = Pt(9)
        label.font.color.rgb = gray
        val = meta.add_run(value)
        val.font.size = Pt(9)
        val.font.color.rgb = gray
    note = doc.add_paragraph()
    legend = [tr_in(options.lang, "exp_uncertain_note_highlight")] if any_marked(t, options) else []
    note_run = note.add_run(sentences(tr_in(options.lang, "exp_generated"), *legend))
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = gray

    colors = {sid: RGBColor.from_string(_SPEAKER_COLORS[i % len(_SPEAKER_COLORS)])
              for i, sid in enumerate(t.speakers)}
    for seg in with_gaps(t):
        if isinstance(seg, Gap):
            marker = doc.add_paragraph().add_run(gap_text(seg, options.lang))
            marker.italic = True
            marker.font.color.rgb = RGBColor(0xC2, 0x41, 0x0C) if seg.speech else gray  # orange: listen here
            continue
        para = doc.add_paragraph()
        para.paragraph_format.keep_together = True
        last = para.add_run(t.speaker_name(seg.speaker))
        last.bold = True
        last.font.color.rgb = colors.get(seg.speaker, gray)
        if options.timestamps:
            last = para.add_run(f"   {clock(seg.start)}")
            last.font.size = Pt(9)
            last.font.color.rgb = gray
        last.add_break()
        for text, uncertain in marked_runs(seg, options):
            run = para.add_run(text)
            if uncertain:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer, tr_in(options.lang, "exp_page"))
    doc.save(str(path))


def _table_header(options: ExportOptions) -> list[str]:
    lang = options.lang
    header = ["#", tr_in(lang, "exp_start"), tr_in(lang, "exp_end"), tr_in(lang, "exp_speaker"), tr_in(lang, "exp_text")]
    return header + [tr_in(lang, "exp_uncertain")] if options.mark_uncertain else header


def _table_row(i: int, row: Row, options: ExportOptions) -> list:
    """Uncertain words get a column of their own, so the text stays as spoken for searching
    and counting, and the rows that need listening can be filtered."""
    cells = [i, clock(row.start, True), clock(row.end, True), _cell(row.speaker), _cell(row.text)]
    if options.mark_uncertain:
        cells.append(_cell(", ".join(row.uncertain)) if row.uncertain else None)
    return cells


_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _cell(text: str) -> str:
    """Text that Excel shows as text: a leading = + - @ would otherwise start a formula
    (spoken words can, and a crafted file name could run one)."""
    return "'" + text if text.startswith(_FORMULA_START) else text


def export_xlsx(t: Transcript, path: Path, options: ExportOptions) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = tr_in(options.lang, "exp_sheet_transcript")
    ws.append(_table_header(options))
    for i, row in enumerate(sentence_rows(t), start=1):
        ws.append(_table_row(i, row, options))
    header_fill = PatternFill("solid", fgColor="E5E7EB")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for col, width in zip("ABCDEF", (6, 14, 14, 16, 100, 28)[:ws.max_column]):
        ws.column_dimensions[col].width = width
    top = Alignment(vertical="top")
    wrap = Alignment(vertical="top", wrap_text=True)
    uncertain_fill = PatternFill("solid", fgColor="FFF59D")  # the yellow of Word's highlight, lighter
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap if cell.column >= 5 else top
            if cell.column == 6 and cell.value:
                cell.fill = uncertain_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet(tr_in(options.lang, "exp_sheet_info"))
    info.append([tr_in(options.lang, "exp_source"), _cell(t.source_name)])
    for key, value in metadata(t, options.lang):
        info.append([key, value])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 80
    for cell in info["A"]:
        cell.font = Font(bold=True)
    wb.save(str(path))


def export_csv(t: Transcript, path: Path, options: ExportOptions) -> None:
    # UTF-8 with BOM so Excel detects the encoding when the file is double-clicked.
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(_table_header(options))
        for i, row in enumerate(sentence_rows(t), start=1):
            writer.writerow(_table_row(i, row, options))
