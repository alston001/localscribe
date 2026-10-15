"""Group D: the repo is public and recordings are private. Nothing personal may leak out."""

import csv
import hashlib
import json
import logging
import os
import shutil
import subprocess
import types
from pathlib import Path

import pytest

from localscribe.core import media, models
from localscribe.core.errors import UserError
from localscribe.core.media import MEDIA_EXTENSIONS
from test_stability import FakeResponse, FakeSession

ROOT = Path(__file__).resolve().parent.parent
needs_git = pytest.mark.skipif(shutil.which("git") is None or not (ROOT / ".git").exists(), reason="needs git")


def ignored(*paths: str) -> set[str]:
    """Which of these (possibly non-existent) paths .gitignore would ignore."""
    # bytes, not text: Windows text mode turns "\n" into "\r\n", which spoils every path but the last
    proc = subprocess.run(["git", "check-ignore", "--no-index", "--stdin"], cwd=ROOT, capture_output=True,
                          input="\n".join(paths).encode())
    return set(proc.stdout.decode().splitlines()) - {""}


# ---- D1: .gitignore ---------------------------------------------------------------
@needs_git
def test_recordings_are_ignored_in_every_folder():
    paths = [f"{folder}rec{ext}" for ext in sorted(MEDIA_EXTENSIONS) for folder in ("", "notes/week 3/")]
    assert ignored(*paths) == set(paths)


@needs_git
def test_transcripts_and_secrets_are_ignored_in_every_folder():
    paths = ["out/class_transcript.docx", "class_transcript (2).json", "a/b/x_transcript.srt",
             "models/whisper/model.bin.part", "logs/localscribe.log", ".env", "sub/.env"]
    assert ignored(*paths) == set(paths)


@needs_git
def test_no_tracked_file_is_ignored():
    proc = subprocess.run(["git", "ls-files", "-ci", "--exclude-standard"], cwd=ROOT, capture_output=True, text=True)
    assert proc.stdout == ""
    keep = ["tests/data/interview_2spk_truth.json", "tests/test_transcript.py", "localscribe/core/transcript.py"]
    assert ignored(*keep) == set()


# ---- D2: no folders in outputs or logs ---------------------------------------------
def test_transcript_stores_the_file_name_only(tmp_path):
    from localscribe.core import pipeline
    from localscribe.core.asr import AsrResult, Word

    words = [Word(0.0, 0.5, " Hello.", 0.9)]
    transcript = pipeline.build_transcript(tmp_path / "private" / "class.m4a", 2.0,
                                           AsrResult("en", 1.0, words), [], pipeline.JobOptions())
    assert transcript.source == "class.m4a"


def test_xlsx_info_sheet_has_no_folder(tmp_path):
    from openpyxl import load_workbook

    from localscribe.export.common import ExportOptions
    from localscribe.export.office import export_xlsx
    from test_export import make_transcript

    t = make_transcript()  # an older transcript that still holds a full path
    export_xlsx(t, tmp_path / "t.xlsx", ExportOptions())
    info = load_workbook(tmp_path / "t.xlsx")["Info"]
    assert info["B1"].value == "interview.mp4"


def test_worker_logs_name_files_not_folders(tmp_path, monkeypatch, caplog):
    from localscribe.export.common import ExportOptions
    from localscribe.ui import threads

    secret = tmp_path / "Clients" / "Jane Doe"
    engine = types.SimpleNamespace(transcribe_file=lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    worker = threads.ProcessThread(engine)
    monkeypatch.setattr("localscribe.core.convert.convert_file", lambda *a, **k: secret / "memo.m4a")
    with caplog.at_level(logging.INFO, logger="localscribe"):
        worker._run(threads.Task(1, secret / "session.wav", None, [], ExportOptions(), None))
        threads.ConvertThread([(0, secret / "memo.qta")], None, "m4a").run()
    assert "session.wav" in caplog.text and "memo.qta" in caplog.text
    assert "Jane Doe" not in caplog.text


# ---- D3: Hugging Face --------------------------------------------------------------
@pytest.mark.parametrize("name", ["../../evil.py", "sub/../../evil.py", "C:/Windows/evil.dll", "/etc/evil"])
def test_repo_files_cannot_escape_the_model_folder(monkeypatch, tmp_path, name):
    listing = [{"type": "file", "path": "config.yaml", "size": 3}, {"type": "file", "path": name, "size": 3}]
    monkeypatch.setattr(models, "_session", lambda token=None: FakeSession(body=json.dumps(listing).encode()))
    started = []
    monkeypatch.setattr(models, "download_files", lambda jobs, **kw: started.append(jobs))
    with pytest.raises(UserError) as caught:
        models.ModelStore(tmp_path).download_pyannote("hf_token")
    assert caught.value.key == "err_unsafe_path" and started == []


def test_token_is_never_sent_over_plain_http(monkeypatch):
    monkeypatch.setattr(models, "_truststore_done", True)
    monkeypatch.setattr(models, "HF_ENDPOINT", "http://mirror.example")
    with pytest.raises(UserError) as caught:
        models._session("hf_secret")
    assert caught.value.key == "err_insecure_endpoint"
    assert "Authorization" not in models._session(None).headers  # no token: plain http is allowed


def test_token_over_https_is_attached(monkeypatch):
    monkeypatch.setattr(models, "_truststore_done", True)
    monkeypatch.setattr(models, "HF_ENDPOINT", "https://huggingface.co")
    assert models._session("hf_secret").headers["Authorization"] == "Bearer hf_secret"


# ---- D4: download integrity --------------------------------------------------------
def test_assets_of_unknown_size_are_refused(monkeypatch, tmp_path):
    monkeypatch.setattr(models, "_remote_size", lambda url: 0)
    started = []
    monkeypatch.setattr(models, "download_files", lambda jobs, **kw: started.append(jobs))
    with pytest.raises(UserError) as caught:
        models.ModelStore(tmp_path).ensure_sherpa()
    assert caught.value.key == "err_size_unknown" and started == []


def test_listing_carries_the_lfs_sha256(monkeypatch):
    digest = hashlib.sha256(b"model").hexdigest()
    listing = [{"type": "file", "path": "model.bin", "size": 1, "lfs": {"oid": digest, "size": 5}},
               {"type": "file", "path": "config.json", "size": 2, "oid": "a" * 40},
               {"type": "directory", "path": "sub"}]
    monkeypatch.setattr(models, "_session", lambda token=None: FakeSession(body=json.dumps(listing).encode()))
    assert models.hf_list_files("org/repo") == [("model.bin", 5, digest), ("config.json", 2, None)]


def test_a_download_with_the_wrong_hash_is_discarded(monkeypatch, tmp_path):
    monkeypatch.setattr(models, "_session", lambda token=None: FakeSession(body=b"evil!"))
    dest = tmp_path / "model.bin"
    with pytest.raises(UserError) as caught:
        models.download_files([("https://example.test/model.bin", dest, 5, hashlib.sha256(b"model").hexdigest())])
    assert caught.value.key == "err_download_corrupt"
    assert list(tmp_path.iterdir()) == []


def test_a_download_with_the_right_hash_is_kept(monkeypatch, tmp_path):
    monkeypatch.setattr(models, "_session", lambda token=None: FakeSession(body=b"model"))
    dest = tmp_path / "model.bin"
    models.download_files([("https://example.test/model.bin", dest, 5, hashlib.sha256(b"model").hexdigest())])
    assert dest.read_bytes() == b"model"


# ---- D5: ffmpeg --------------------------------------------------------------------
def capture_commands(monkeypatch):
    seen = []

    def run(cmd, **kwargs):
        seen.append(cmd)
        return types.SimpleNamespace(returncode=0, stderr=b"  Duration: 00:00:02.00, start: 0\n"
                                     b"  Stream #0:0: Audio: pcm_s16le, 16000 Hz, mono, s16\n")

    class Popen:
        def __init__(self, cmd, **kwargs):
            seen.append(cmd)
            raise OSError("stop here")

    monkeypatch.setattr(media, "ffmpeg_exe", lambda: "ffmpeg")
    monkeypatch.setattr(media.subprocess, "run", run)
    monkeypatch.setattr(media.subprocess, "Popen", Popen)
    return seen


def test_ffmpeg_may_only_read_local_files(monkeypatch, tmp_path):
    from localscribe.core import convert

    seen = capture_commands(monkeypatch)
    monkeypatch.setattr(convert, "ffmpeg_exe", lambda: "ffmpeg", raising=False)
    media.probe(tmp_path / "a.wav")
    for call in (lambda: media.load_audio(tmp_path / "a.wav"),
                 lambda: convert.convert_file(tmp_path / "a.wav", tmp_path, "mp3")):
        try:
            call()
        except Exception:
            pass
    kinds = {"probe": False, "decode": False, "convert": False}
    for cmd in seen:
        kinds["convert" if "-progress" in cmd else "decode" if "pipe:1" in cmd else "probe"] = True
    assert all(kinds.values()), kinds
    for cmd in seen:
        assert "-protocol_whitelist" in cmd, cmd
        at = cmd.index("-protocol_whitelist")
        assert cmd[at + 1] == "file,pipe" and at < cmd.index("-i")


def test_metadata_cannot_fake_the_duration_or_streams(tmp_path):
    from test_media import ffmpeg

    tone = tmp_path / "tone.wav"
    ffmpeg("-f", "lavfi", "-i", "sine=frequency=440:duration=2", "-ar", "16000", str(tone))
    tricky = tmp_path / "tricky.mka"
    ffmpeg("-i", str(tone), "-metadata", "title=Duration: 05:00:00.00, Stream #0:7: Audio: none (apac / x), 4.0",
           "-c:a", "flac", str(tricky))
    info = media.probe(tricky)
    assert abs(info.duration - 2.0) < 0.2
    assert len(info.audio) == 1 and info.audio[0].decodable


# ---- D6: spreadsheet formulas ------------------------------------------------------
def risky_transcript():
    from test_export import make_transcript

    t = make_transcript()
    t.speakers["S1"] = "@admin"
    for seg, text in zip(t.segments, ["=HYPERLINK(\"http://x\",\"y\")", "+1 for that.", "-5 degrees."]):
        seg.text, seg.words = text, []
    return t


def test_csv_cells_cannot_start_a_formula(tmp_path):
    from localscribe.export.common import ExportOptions
    from localscribe.export.office import export_csv

    export_csv(risky_transcript(), tmp_path / "t.csv", ExportOptions())
    with open(tmp_path / "t.csv", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))[1:]
    assert [r[3] for r in rows] == ["'@admin", "Speaker 2", "'@admin"]
    assert [r[4] for r in rows] == ["'=HYPERLINK(\"http://x\",\"y\")", "'+1 for that.", "'-5 degrees."]
    assert rows[0][1] == "00:00:00.000"  # timestamps are untouched


def test_xlsx_cells_cannot_start_a_formula(tmp_path):
    from openpyxl import load_workbook

    from localscribe.export.common import ExportOptions
    from localscribe.export.office import export_xlsx

    export_xlsx(risky_transcript(), tmp_path / "t.xlsx", ExportOptions())
    ws = load_workbook(tmp_path / "t.xlsx").worksheets[0]
    cells = [ws.cell(row, 5) for row in range(2, 5)]
    assert [c.value for c in cells] == ["'=HYPERLINK(\"http://x\",\"y\")", "'+1 for that.", "'-5 degrees."]
    assert all(c.data_type == "s" for c in cells) and ws.cell(2, 4).value == "'@admin"


# ---- D7: explorer ------------------------------------------------------------------
def test_explorer_is_started_by_its_full_path(monkeypatch, tmp_path):
    from localscribe.ui import convert_dialog, main_window

    monkeypatch.setenv("WINDIR", r"C:\Windows")
    started = []
    for module in (main_window, convert_dialog):
        monkeypatch.setattr(module.subprocess, "Popen", lambda cmd, **k: started.append(cmd))
    target = str(tmp_path / "out.docx")
    main_window.MainWindow._open_folder(None, types.SimpleNamespace(outputs=[target], path=Path("x")))
    item = types.SimpleNamespace(data=lambda role: target)
    convert_dialog.ConvertDialog._reveal(None, item)
    assert [cmd[0] for cmd in started] == [os.path.join(r"C:\Windows", "explorer.exe")] * 2
    assert all(cmd[-1] == os.path.normpath(target) for cmd in started)
