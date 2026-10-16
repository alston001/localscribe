import csv
import json

from localscribe import i18n
from localscribe.config import Settings
from localscribe.core.transcript import Gap, Segment, Transcript, TWord
from localscribe.export import FORMATS, output_base, write_outputs
from localscribe.export.common import (UNCERTAIN_BELOW, ExportOptions, clock, marked_runs, marked_text,
                                       sentence_rows, subtitle_cues)
from localscribe.export.office import export_csv, export_docx, export_xlsx
from localscribe.export.text import export_md, export_txt


def make_transcript(language="en"):
    def seg(speaker, start, text):
        tokens, t, ws = text.split(), start, []
        for tok in tokens:
            ws.append(TWord(t, t + 0.3, " " + tok, 0.9))
            t += 0.35
        return Segment(speaker, ws[0].start, ws[-1].end, text, ws)

    return Transcript(
        source="C:/rec/interview.mp4", duration=30.0, language=language, created="2026-10-15T10:00:00",
        app_version="test", settings={"whisper_model": "large-v3", "diarizer": "pyannote"},
        speakers={"S1": "Speaker 1", "S2": "Speaker 2"},
        segments=[
            seg("S1", 0.0, "Thank you for coming. Could you introduce yourself?"),
            seg("S2", 5.0, "Sure. I study how older adults use voice assistants at home every day."),
            seg("S1", 12.0, "Great."),
        ])


def test_clock_formats():
    assert clock(3723.456) == "01:02:03"
    assert clock(3723.456, millis=True, sep=",") == "01:02:03,456"
    assert clock(-1) == "00:00:00"


def test_sentence_rows_split_segments():
    rows = sentence_rows(make_transcript())
    assert [r[3] for r in rows][:3] == ["Thank you for coming.", "Could you introduce yourself?", "Sure."]
    assert rows[2][2] == "Speaker 2"


def test_subtitles_label_speaker_changes_only_and_avoid_orphans():
    cues = subtitle_cues(make_transcript(), ExportOptions())
    labelled = [c for c in cues if c.speaker]
    assert [c.speaker for c in labelled] == ["Speaker 1", "Speaker 2", "Speaker 1"]
    texts = [" ".join(c.lines) for c in cues]
    assert not any(t.strip() in ("home.", "day.") for t in texts)
    assert all(len(line) <= 60 for c in cues for line in c.lines)
    for a, b in zip(cues, cues[1:]):
        assert a.end <= b.start


def test_output_base_never_overwrites(tmp_path):
    (tmp_path / "talk_transcript.docx").write_text("x")
    assert output_base(tmp_path, "talk", ["txt"]) == "talk_transcript"
    assert output_base(tmp_path, "talk", ["docx", "txt"]) == "talk_transcript (2)"


def test_write_all_formats(tmp_path):
    written = write_outputs(make_transcript(), tmp_path, list(FORMATS), ExportOptions(lang="zh_TW"))
    assert [p.suffix for p in written] == [".json"] + [ext for ext, _ in FORMATS.values()]
    data = json.loads(written[0].read_text(encoding="utf-8"))
    assert data["format_version"] == 1 and data["speakers"]["S2"] == "Speaker 2"
    assert (tmp_path / "interview_transcript.txt").read_text(encoding="utf-8").count("Speaker 1:") == 2
    assert (tmp_path / "interview_transcript.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    srt = (tmp_path / "interview_transcript.srt").read_text(encoding="utf-8")
    assert srt.startswith("1\n00:00:00,000 --> ")


def test_transcript_round_trip(tmp_path):
    t = make_transcript()
    t.save(tmp_path / "t.json")
    again = Transcript.load(tmp_path / "t.json")
    assert again.segments[1].words[0].text == " Sure."
    assert again.speaker_name("S1") == "Speaker 1"


def test_settings_ignore_bad_values(tmp_path):
    path = tmp_path / "settings.json"
    path.write_text(json.dumps({"formats": ["docx", "bogus"], "speakers_exact": "two", "preset": "zzz",
                                "unknown_key": 1}), encoding="utf-8")
    s = Settings.load(path)
    assert s.formats == ["docx"] and s.speakers_exact == 2 and s.preset == "custom"
    path.write_text("{not json", encoding="utf-8")
    assert Settings.load(path).formats == ["docx", "txt"]


def test_settings_from_older_version_get_current_preset_values(tmp_path):
    path = tmp_path / "settings.json"
    path.write_text(json.dumps({"preset": "lecture", "speaker_mode": "range", "speakers_min": 1,
                                "speakers_max": 5}), encoding="utf-8")
    s = Settings.load(path)
    assert s.speaker_mode == "auto" and s.speaker_count().max == 0
    path.write_text(json.dumps({"preset": "custom", "speaker_mode": "range", "speakers_min": 1,
                                "speakers_max": 5}), encoding="utf-8")
    custom = Settings.load(path)
    assert (custom.speaker_mode, custom.speakers_max) == ("range", 5)  # hand-picked values stay


# ---- uncertain words ---------------------------------------------------------------
LOW, HIGH = UNCERTAIN_BELOW / 2, 0.95
MARK = ExportOptions(mark_uncertain=True)


def unsure_transcript():
    """Whisper was unsure of the two terms; "Right." it heard clearly."""
    words = [TWord(0.0, 0.3, " We", HIGH), TWord(0.3, 0.6, " start", HIGH), TWord(0.6, 0.9, " with", HIGH),
             TWord(0.9, 1.5, " tokenization,", LOW), TWord(1.6, 1.9, " then", HIGH),
             TWord(1.9, 2.6, " lemmatization.", LOW)]
    return Transcript(
        source="class.wav", duration=4.0, language="en", created="2026-10-15T10:00:00", app_version="test",
        settings={}, speakers={"S1": "Speaker 1", "S2": "Speaker 2"},
        segments=[Segment("S1", 0.0, 2.6, "We start with tokenization, then lemmatization.", words),
                  Segment("S2", 3.0, 3.5, "Right.", [TWord(3.0, 3.5, " Right.", HIGH)])])


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.reader(fh))


def test_uncertain_words_get_a_mark_in_txt_and_md(tmp_path):
    t = unsure_transcript()
    export_txt(t, tmp_path / "a.txt", MARK)
    export_md(t, tmp_path / "a.md", MARK)
    for name in ("a.txt", "a.md"):
        text = (tmp_path / name).read_text(encoding="utf-8")
        assert "We start with tokenization(?), then lemmatization(?)." in text, name
        assert "Right." in text and "Right(?)" not in text


def test_uncertain_words_are_highlighted_in_word(tmp_path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    export_docx(unsure_transcript(), tmp_path / "a.docx", MARK)
    paragraphs = Document(str(tmp_path / "a.docx")).paragraphs
    highlighted = [r.text for p in paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert highlighted == ["tokenization", "lemmatization"]  # the word, not its comma
    assert any(p.text.endswith("\nWe start with tokenization, then lemmatization.") for p in paragraphs)


def test_tables_list_uncertain_words_in_their_own_column(tmp_path):
    # The text column stays as spoken, for searching and counting; the extra column can be filtered.
    from openpyxl import load_workbook

    t = unsure_transcript()
    export_csv(t, tmp_path / "a.csv", MARK)
    rows = read_csv(tmp_path / "a.csv")
    assert rows[0][-1] == "Unsure words"
    assert [r[4:] for r in rows[1:]] == [["We start with tokenization, then lemmatization.", "tokenization, lemmatization"],
                                         ["Right.", ""]]
    export_xlsx(t, tmp_path / "a.xlsx", ExportOptions(lang="zh_TW", mark_uncertain=True))
    sheet = load_workbook(tmp_path / "a.xlsx").worksheets[0]
    assert [c.value for c in sheet[1]][-1] == "不確定的字詞"
    assert sheet["F2"].value == "tokenization, lemmatization" and sheet["F3"].value is None


def test_an_uncertain_word_cannot_start_a_formula(tmp_path):
    t = unsure_transcript()
    t.segments[1].words[0] = TWord(3.0, 3.5, " +1.", LOW)
    t.segments[1].text = "+1."
    export_csv(t, tmp_path / "a.csv", MARK)
    assert read_csv(tmp_path / "a.csv")[2][4:] == ["'+1.", "'+1"]


def test_subtitles_are_never_marked(tmp_path):
    t = unsure_transcript()
    for fmt in ("srt", "vtt"):
        writer = FORMATS[fmt][1]
        writer(t, tmp_path / f"plain.{fmt}", ExportOptions())
        writer(t, tmp_path / f"marked.{fmt}", MARK)
        assert (tmp_path / f"plain.{fmt}").read_text(encoding="utf-8") == \
            (tmp_path / f"marked.{fmt}").read_text(encoding="utf-8")


def test_nothing_is_marked_unless_asked(tmp_path):
    from docx import Document

    write_outputs(unsure_transcript(), tmp_path, list(FORMATS), ExportOptions())
    for ext in ("txt", "md", "srt", "vtt", "csv"):
        assert "(?)" not in (tmp_path / f"class_transcript.{ext}").read_text(encoding="utf-8-sig"), ext
    paragraphs = Document(str(tmp_path / "class_transcript.docx")).paragraphs
    assert not any(r.font.highlight_color for p in paragraphs for r in p.runs)
    assert len(read_csv(tmp_path / "class_transcript.csv")[0]) == 5


def test_edited_text_has_no_confidences_to_mark(tmp_path):
    t = unsure_transcript()
    t.segments[0].text = "We start with tokenization, then lemmatisation."  # corrected; the words were not
    export_txt(t, tmp_path / "a.txt", MARK)
    export_csv(t, tmp_path / "a.csv", MARK)
    lines = (tmp_path / "a.txt").read_text(encoding="utf-8").splitlines()
    assert "[00:00:00] Speaker 1: We start with tokenization, then lemmatisation." in lines
    assert read_csv(tmp_path / "a.csv")[1][5] == ""


def test_the_mark_goes_after_the_word_and_before_its_punctuation():
    seg = Segment("S1", 0.0, 2.0, '"Tokenization," she said.',
                  [TWord(0.0, 1.0, ' "Tokenization,"', LOW), TWord(1.0, 1.5, " she", HIGH),
                   TWord(1.5, 2.0, " said.", HIGH)])
    assert marked_text(seg, MARK) == '"Tokenization(?)," she said.'
    assert marked_runs(seg, MARK) == [('"', False), ("Tokenization", True), ('," she said.', False)]


def test_cjk_words_are_marked_without_adding_spaces():
    seg = Segment("S1", 0.0, 1.0, "這是分詞。",
                  [TWord(0.0, 0.3, "這是", HIGH), TWord(0.3, 0.6, "分詞", LOW), TWord(0.6, 1.0, "。", LOW)])
    assert marked_text(seg, MARK) == "這是分詞(?)。"  # a lone punctuation mark is never marked


def test_words_at_the_threshold_are_not_marked():
    seg = Segment("S1", 0.0, 1.0, "a b",
                  [TWord(0.0, 0.5, " a", UNCERTAIN_BELOW), TWord(0.5, 1.0, " b", UNCERTAIN_BELOW - 0.001)])
    assert marked_text(seg, MARK) == "a b(?)"


def test_without_marking_the_text_is_returned_as_it_is():
    seg = unsure_transcript().segments[0]
    assert marked_text(seg, ExportOptions()) == seg.text
    assert marked_runs(seg, ExportOptions()) == [(seg.text, False)]


def test_marking_uncertain_words_is_off_by_default_and_remembered(tmp_path):
    assert ExportOptions().mark_uncertain is False and Settings().mark_uncertain is False
    Settings(mark_uncertain=True).save(tmp_path / "settings.json")
    assert Settings.load(tmp_path / "settings.json").export_options("en").mark_uncertain is True


def test_uncertain_word_strings_have_both_languages():
    for key in ("mark_uncertain", "mark_uncertain_tip", "exp_uncertain", "exp_uncertain_note_mark",
                "exp_uncertain_note_highlight"):
        assert key in i18n.STRINGS and all(i18n.STRINGS[key]), key


def test_command_line_can_mark_uncertain_words(monkeypatch, tmp_path):
    from test_vocabulary import FakeEngine

    from localscribe import cli

    seen = []
    monkeypatch.setattr(cli, "Engine", FakeEngine)
    monkeypatch.setattr(cli, "write_outputs", lambda t, folder, formats, options: seen.append(options) or [])
    for flags in ([], ["--mark-uncertain"]):
        assert cli.main([str(tmp_path / "talk.wav"), "--models-dir", str(tmp_path), *flags]) == 0
    assert [o.mark_uncertain for o in seen] == [False, True]



# ---- gaps: long stretches without text ----------------------------------------------
def gappy_transcript():
    def seg(speaker, start, end, text):
        return Segment(speaker, start, end, text, [TWord(start, end, " " + text, 0.9)])

    return Transcript(
        source="lecture.wav", duration=150.0, language="en", created="2026-10-15T10:00:00", app_version="test",
        settings={}, speakers={"S1": "Speaker 1", "S2": "Speaker 2"},
        segments=[seg("S1", 0.0, 2.0, "Welcome back."), seg("S1", 45.0, 47.0, "Let's continue."),
                  seg("S2", 140.0, 142.0, "One question?")],
        gaps=[Gap(2.0, 45.0), Gap(47.0, 100.0, speech=True), Gap(100.0, 140.0)])


def body(path) -> list[str]:
    """The non-empty lines after the heading block."""
    lines = path.read_text(encoding="utf-8").splitlines()
    return [line for line in lines[lines.index("") + 1:] if line and line != "---"]


def test_gaps_are_marked_between_paragraphs_in_txt(tmp_path):
    export_txt(gappy_transcript(), tmp_path / "a.txt", ExportOptions())
    assert body(tmp_path / "a.txt") == [
        "[00:00:00] Speaker 1: Welcome back.",
        "[Nothing transcribed 00:00:02–00:00:45]",
        "[00:00:45] Speaker 1: Let's continue.",
        "[Speech detected, nothing transcribed 00:00:47–00:01:40]",
        "[Nothing transcribed 00:01:40–00:02:20]",
        "[00:02:20] Speaker 2: One question?"]


def test_gap_markers_follow_the_export_language(tmp_path):
    export_txt(gappy_transcript(), tmp_path / "a.txt", ExportOptions(lang="zh_TW", timestamps=False))
    lines = body(tmp_path / "a.txt")
    assert lines[1] == "[沒有轉出文字 00:00:02–00:00:45]"  # the times stay: they say where to listen
    assert lines[3] == "[偵測到語音，但沒有轉出文字 00:00:47–00:01:40]"


def test_gaps_are_marked_in_md_and_docx(tmp_path):
    from docx import Document

    t = gappy_transcript()
    export_md(t, tmp_path / "a.md", ExportOptions())
    md = (tmp_path / "a.md").read_text(encoding="utf-8")
    assert md.index("Welcome back.") < md.index("*[Nothing transcribed 00:00:02–00:00:45]*") < md.index("Let's continue.")
    export_docx(t, tmp_path / "a.docx", ExportOptions())
    texts = [p.text for p in Document(str(tmp_path / "a.docx")).paragraphs]
    start = next(i for i, text in enumerate(texts) if text.endswith("Welcome back."))
    assert texts[start + 1:start + 5] == ["[Nothing transcribed 00:00:02–00:00:45]", texts[start + 2],
                                          "[Speech detected, nothing transcribed 00:00:47–00:01:40]",
                                          "[Nothing transcribed 00:01:40–00:02:20]"]
    assert texts[start + 2].endswith("Let's continue.")


def test_a_gap_filled_in_later_is_not_marked(tmp_path):
    t = gappy_transcript()
    t.segments.insert(2, Segment("S1", 60.0, 62.0, "(inaudible question)"))  # typed in by an editor
    export_txt(t, tmp_path / "a.txt", ExportOptions())
    lines = body(tmp_path / "a.txt")
    assert not any("Speech detected" in line for line in lines) and "[Nothing transcribed 00:01:40–00:02:20]" in lines


def test_subtitles_and_tables_have_no_gap_markers(tmp_path):
    # Subtitles show silence by their timing, and a marker row would not fit a table of sentences.
    write_outputs(gappy_transcript(), tmp_path, list(FORMATS), ExportOptions())
    for ext in ("srt", "vtt", "csv"):
        text = (tmp_path / f"lecture_transcript.{ext}").read_text(encoding="utf-8-sig")
        assert "Nothing transcribed" not in text and "Speech detected" not in text, ext


def test_gap_strings_have_both_languages():
    for key in ("exp_gap_silence", "exp_gap_speech"):
        assert key in i18n.STRINGS and all(i18n.STRINGS[key]), key


def test_the_legend_explains_the_marks_in_the_headings_language(tmp_path):
    from docx import Document

    zh = ExportOptions(lang="zh_TW", mark_uncertain=True)
    export_md(unsure_transcript(), tmp_path / "a.md", zh)
    export_docx(unsure_transcript(), tmp_path / "a.docx", zh)
    assert "*由 LocalScribe 在本機離線產生，請在使用前校對。Whisper 沒把握的字後面標有 (?)。*" in \
        (tmp_path / "a.md").read_text(encoding="utf-8").splitlines()  # no space after a Chinese full stop
    texts = [p.text for p in Document(str(tmp_path / "a.docx")).paragraphs]
    assert "由 LocalScribe 在本機離線產生，請在使用前校對。Whisper 沒把握的字以黃色標出。" in texts
    export_md(unsure_transcript(), tmp_path / "b.md", MARK)
    assert "*Generated offline by LocalScribe. Please proofread before use. (?) follows words Whisper was " \
           "unsure of.*" in (tmp_path / "b.md").read_text(encoding="utf-8").splitlines()



def test_a_plain_gap_does_not_claim_silence():
    # The voice detector may have heard up to 30 s of speech there (a dropped question), so the
    # marker says only that nothing was written; "Speech detected" is kept for the lost-speech stretches.
    for lang in ("en", "zh_TW"):
        plain = i18n.tr_in(lang, "exp_gap_silence", start="0", end="1")
        assert "speech" not in plain.lower() and "語音" not in plain


def test_marking_never_changes_the_text_s_spacing():
    # Words match a text up to whitespace; an edited line break must survive the marking.
    seg = Segment("S1", 0.0, 1.0, "Hello\nworld.", [TWord(0.0, 0.5, " Hello", HIGH), TWord(0.5, 1.0, " world.", LOW)])
    assert marked_text(seg, MARK) == "Hello\nworld."


def test_whitespace_only_and_glued_words_keep_the_text():
    seg = Segment("S1", 0.0, 1.0, "skip-gram model",
                  [TWord(0.0, 0.2, " skip", LOW), TWord(0.2, 0.4, "-gram", LOW), TWord(0.4, 0.5, " ", LOW),
                   TWord(0.5, 1.0, "model", HIGH)])
    assert marked_text(seg, MARK) == "skip-gram(?) model"  # glued pieces form one mark


def test_marks_respect_abbreviations_quotes_and_dashes():
    def mark(text):
        return marked_text(Segment("S1", 0.0, 1.0, text, [TWord(0.0, 1.0, " " + text, LOW)]), MARK)

    assert mark("U.S.") == "U.S.(?)"
    assert mark("e.g.,") == "e.g.(?),"
    assert mark("students'") == "students'(?)"
    assert mark("'this'") == "'this(?)'"
    assert mark("„Wort“") == "„Wort(?)“"
    assert mark("「分詞」") == "「分詞(?)」"
    assert mark("word—") == "word(?)—"


def test_adjacent_unsure_cjk_characters_form_one_mark(tmp_path):
    words = [TWord(0.0, 0.2, "我", LOW), TWord(0.2, 0.4, "想", LOW), TWord(0.4, 0.6, "去", LOW),
             TWord(0.6, 1.0, "學校。", HIGH)]
    t = unsure_transcript()
    t.segments = [Segment("S1", 0.0, 1.0, "我想去學校。", words)]
    assert marked_text(t.segments[0], MARK) == "我想去(?)學校。"
    export_csv(t, tmp_path / "a.csv", MARK)
    assert read_csv(tmp_path / "a.csv")[1][5] == "我想去"


def test_a_word_is_judged_by_the_probability_the_json_keeps(tmp_path):
    # _transcript.json keeps 3 decimals; a re-export from it must mark the same words.
    t = unsure_transcript()
    t.segments[1].words[0].prob = UNCERTAIN_BELOW - 0.0004
    export_txt(t, tmp_path / "a.txt", MARK)
    t.save(tmp_path / "t.json")
    export_txt(Transcript.load(tmp_path / "t.json"), tmp_path / "b.txt", MARK)
    assert (tmp_path / "a.txt").read_text(encoding="utf-8") == (tmp_path / "b.txt").read_text(encoding="utf-8")


def test_an_edited_segment_has_no_highlights_in_word(tmp_path):
    from docx import Document

    t = unsure_transcript()
    t.segments[0].text = "We start with tokenization, then lemmatisation."
    export_docx(t, tmp_path / "a.docx", MARK)
    paragraphs = Document(str(tmp_path / "a.docx")).paragraphs
    assert not any(r.font.highlight_color for p in paragraphs for r in p.runs)
    assert any(p.text.endswith("\nWe start with tokenization, then lemmatisation.") for p in paragraphs)


def test_gaps_before_after_and_out_of_order():
    from localscribe.export.common import with_gaps

    t = gappy_transcript()
    t.gaps = [Gap(142.0, 180.0), Gap(100.0, 140.0), Gap(-1.0, 0.0), Gap(47.0, 100.0, speech=True),
              Gap(44.0, 46.0)]  # unsorted; the last overlaps "Let's continue." partly
    t.duration = 180.0
    order = [(type(item).__name__, item.start) for item in with_gaps(t)]
    assert order == [("Gap", -1.0), ("Segment", 0.0), ("Segment", 45.0), ("Gap", 47.0), ("Gap", 100.0),
                     ("Segment", 140.0), ("Gap", 142.0)]


def test_the_checkbox_is_saved_and_reaches_the_exports(tmp_path):
    # Real widgets, offscreen, in their own process; the hardware probe is replaced so nothing loads torch.
    import os
    import subprocess
    import sys
    from pathlib import Path

    script = r'''
import json, sys
from pathlib import Path
from localscribe import paths
paths.user_data_dir = lambda: Path(sys.argv[1])
from PySide6.QtCore import QThread, Signal
from PySide6.QtWidgets import QApplication
app = QApplication([])
from localscribe.config import Settings
from localscribe.ui import main_window

class NoProbe(QThread):
    ready = Signal(object)

    def run(self):
        pass

main_window.DeviceProbe = NoProbe
window = main_window.MainWindow(Settings())
assert not window.chk_uncertain.isChecked()
window.chk_uncertain.setChecked(True)
saved = json.loads((Path(sys.argv[1]) / "settings.json").read_text(encoding="utf-8"))
again = main_window.MainWindow(Settings.load(Path(sys.argv[1]) / "settings.json"))
print(json.dumps([saved["mark_uncertain"], again.chk_uncertain.isChecked(),
                  again.settings.export_options("en").mark_uncertain]))
for w in (window, again):
    for thread in list(w._helpers):
        thread.wait()
'''
    env = {**os.environ, "QT_QPA_PLATFORM": "offscreen"}
    root = Path(__file__).resolve().parent.parent
    proc = subprocess.run([sys.executable, "-c", script, str(tmp_path)], cwd=root, env=env,
                          capture_output=True, text=True, timeout=60)
    assert proc.returncode == 0, proc.stderr
    assert json.loads(proc.stdout.strip().splitlines()[-1]) == [True, True, True]
